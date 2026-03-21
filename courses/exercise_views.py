
from django.views.generic import ListView, CreateView, UpdateView, DetailView, View
from rest_framework.views import APIView
from rest_framework.decorators import api_view
from django.urls import reverse, reverse_lazy
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from openpyxl import Workbook, load_workbook
from .models import Exercise, ExerciseSet,Chapter,ExerciseSetCompletion,ExerciseAnswer,ExerciseSetCompletion,KnowledgePoint,ExerciseKnowledgePoint
from .forms import ExerciseForm
from .serializers import ExerciseSerializer, ExerciseSetCompletionSerializer
from rest_framework.pagination import PageNumberPagination
from rest_framework import generics, status
from rest_framework.response import Response
from django.shortcuts import get_object_or_404
from django.utils import timezone
import json
from django.utils import timezone
from rest_framework.response import Response

class StandardResultsSetPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = 'page_size'
    max_page_size = 100

# 习题列表视图
class ExerciseListView(ListView):
    serializer_class = ExerciseSerializer
    queryset = Exercise.objects.all().order_by('-updated_at')
    template_name = 'courses/_exercise.html'

    def get(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        serializer = self.serializer_class(queryset, many=True)
        data = serializer.data
        
        # 自定义分页实现
        page_size = int(request.GET.get('page_size', 10))
        page = int(request.GET.get('page', 1))
        
        start = (page - 1) * page_size
        end = start + page_size
        paginated_data = data[start:end]
        
        return JsonResponse({
            'results': paginated_data,
            'count': len(data),
            'current_page': page,
            'page_size': page_size,
            'total_pages': (len(data) + page_size - 1) // page_size
        })

    def get_queryset(self):
        queryset = super().get_queryset()
        title = self.request.GET.get('title', '')
        update_date_start = self.request.GET.get('update_date_start', '')
        update_date_end = self.request.GET.get('update_date_end', '')
        chapter_id = self.request.GET.get('chapter_id', '')
        
        # 过滤掉有parent_id的子题
        queryset = queryset.filter(parent_id__isnull=True)
        
        if title:
            queryset = queryset.filter(title__icontains=title)
        if update_date_start and update_date_end:
            queryset = queryset.filter(updated_at__range=[update_date_start, update_date_end])
        if chapter_id:
            queryset = queryset.filter(chapter__id=chapter_id)
            
        return queryset

    def get_paginated_response(self, data):
        response = super().get_paginated_response(data)
        if not data:
            response.data['results'] = []
            response.data['empty_message'] = '暂无数据'
            response.data['can_create'] = True
        # 添加前端需要的分页字段
        response.data['current_page'] = self.paginator.page.number
        response.data['count'] = self.paginator.page.paginator.count
        return response


class ExerciseCreateView(View):    
    def post(self, request):
        form = ExerciseForm(request.POST)
        if form.is_valid():
            exercise = form.save()
            
            # 处理知识点关联
            tag_data_str = request.POST.get('tag_data', '[]')
            try:
                tag_data = json.loads(tag_data_str)
                self.process_knowledge_points(exercise, tag_data)
            except json.JSONDecodeError:
                return JsonResponse({
                    'success': False,
                    'errors': '知识点数据格式错误'
                }, status=400)
            
            return JsonResponse({
                'success': True,
                'message': '习题创建成功',
                'exercise_id': exercise.id
            })
        return JsonResponse({
            'success': False,
            'errors': form.errors
        }, status=400)
    
    def process_knowledge_points(self, exercise, tag_data):
        """
        处理知识点关联--
        :param exercise: 习题对象
        :param tag_data: 知识点数据列表，格式:
            [
                { "title": "知识点1", "id": "123" }
            ]
        """
        from django.db import transaction
        
        with transaction.atomic():
            # 清除现有关联
            ExerciseKnowledgePoint.objects.filter(exercise=exercise).delete()
            
            # 创建新的关联
            for tag in tag_data:
                title = tag.get('title')
                point_id = tag.get('id')
                
                if not title:
                    continue  # 必须有标题
                
                # 查找或创建知识点对象
                if point_id:
                    try:
                        knowledge_point = KnowledgePoint.objects.get(id=point_id)
                    except KnowledgePoint.DoesNotExist:
                        # 如果ID存在但找不到对象，可以记录日志或抛出错误
                        continue
                else:
                    # 如果没有ID，尝试根据标题查找
                    knowledge_point, created = KnowledgePoint.objects.get_or_create(
                        title=title,
                        defaults={'created_by': self.request.user}  # 可以根据需要设置默认值
                    )
                
                # 创建关联
                ExerciseKnowledgePoint.objects.create(
                    exercise=exercise,
                    knowledge_point=knowledge_point
                )

class ExerciseImportView(View):
    def get(self, request):
        # 提供模板下载
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=exercise_template.xlsx'
        
        # 创建Excel模板
        wb = Workbook()
        ws = wb.active
        ws.append(['标题','题目', '题型', '答案', '解析', '难度', '章节ID','父题目ID','题目顺序','选项'])
        ws.append(['示例标题', '示例题目', 'single_choice', 'A', '解析内容', 3, 1, 10,1,"json格式"])
        wb.save(response)
        return response
    
    def post(self, request):
        try:
            file = request.FILES['file']
            # 解析Excel文件
            wb = load_workbook(filename=file)
            ws = wb.active
            created_count = 0
            skipped_count = 0
            errors = []
            
            # 获取有效的question_type choices
            valid_question_types = dict(Exercise.QUESTION_TYPE_CHOICES).keys()
            
            for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                # 校验question_type
                if row[2] not in valid_question_types:
                    skipped_count += 1
                    errors.append(f"第{idx}行: 无效的题型 '{row[2]}'，有效值为 {', '.join(valid_question_types)}")
                    continue
                    
                try:
                    # 统一处理options字段格式
                    options = row[9]
                    if isinstance(options, str):
                        try:
                            # 处理转义JSON字符串的情况
                            if '\\"' in options:
                                options = json.loads(options)
                            else:
                                # 尝试直接解析为JSON
                                options = json.loads(options) if options.strip().startswith('{') else options
                        except json.JSONDecodeError:
                            # 如果解析失败，保持原样
                            pass
                    
                    Exercise.objects.create(
                        title=row[0],
                        content=row[1],
                        question_type=row[2],
                        answer=row[3],
                        analysis=row[4],
                        difficulty=row[5],
                        chapter_id=row[6],
                        parent_id=row[7] if row[7] else None,
                        order=row[8],
                        options=options if isinstance(options, dict) else {"raw_options": str(options)}
                    )
                    created_count += 1
                except Exception as e:
                    skipped_count += 1
                    errors.append(f"第{idx}行: 创建失败 - {str(e)}")
                
            result = {
                'success': True,
                'created_count': created_count,
                'skipped_count': skipped_count,
                'message': f'成功导入 {created_count} 条习题，跳过 {skipped_count} 条'
            }
            if errors:
                result['errors'] = errors[:10]  # 最多返回10条错误
            return JsonResponse(result)
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'导入失败: {str(e)}'
            }, status=400)

class ExerciseUpdateView(UpdateView):
    model = Exercise
    form_class = ExerciseForm
    template_name = 'courses/exercise_form.html'
    
    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        view_mode = request.GET.get('view_mode', '0') == '1'
        return self.render_to_response(self.get_context_data(
            form=self.get_form(),
            view_mode=view_mode
        ))
    
    def get_success_url(self):
        messages.success(self.request, '习题更新成功')
        return reverse_lazy('courses:exercise_list')

class ExerciseDetailView(DetailView):
    """习题详情视图"""
    model = Exercise
    template_name = 'courses/exercise_detail.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['user_answers'] = self.object.answers.filter(
            user=self.request.user
        ).order_by('-created_at')[:5]
        return context

class ExerciseEditView(UpdateView):
    """习题编辑视图"""
    model = Exercise
    form_class = ExerciseForm
    template_name = 'courses/exercise_edit.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['exercise'] = self.get_object()
        return context
        
    def form_valid(self, form):
        response = super().form_valid(form)
        if self.request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'message': '习题更新成功',
                'exercise_id': self.object.id
            })
        return response

class ExerciseViewView(DetailView):
    """习题查看视图"""
    model = Exercise
    template_name = 'courses/exercise_view.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        exercise = self.get_object()
        
        # 预加载关联知识点
        exercise = Exercise.objects.prefetch_related(
            'knowledge_points'
        ).get(id=exercise.id)
        
        # 获取知识点数据
        knowledge_points = [
            {
                'id': kp.id,
                'title': kp.title
            }
            for kp in exercise.knowledge_points.all()
        ]

        # 获取习题关联的方法总结
        methods = [
            {
                'id': m.id,
                'title': m.name
            }
            for m in exercise.methodsummary_set.all()
        ]

        # 获取带习题集信息的答题记录
        enhanced_records = ExerciseAnswer.objects.filter(
            exercise=exercise
        ).select_related(
            'exercise_set_completion__exercise_set'
        ).order_by('-created_at')[:50]
        
        answer_records = []
        for record in enhanced_records:
            record_data = {
                'exercise': record.exercise,
                'answer': record.answer,
                'is_correct': record.is_correct,
                'spent_time': record.time_spent,
                'creat_at': record.created_at
            }
            
            # 统一处理关联数据
            if record.exercise_set_completion:
                record_data.update({
                    'exerciseset': record.exercise_set_completion.exercise_set.name,
                    'exerciseset_id': record.exercise_set_completion.exercise_set_id,
                    'is_from_set': True
                })
            else:
                record_data.update({
                    'exerciseset': None,
                    'exerciseset_id': 1,
                    'is_from_set': False
                })
            
            answer_records.append(record_data)

        # 如果是完形填空或阅读理解，获取子题
        sub_exercises = []
        if exercise.question_type in ['cloze_test', 'reading_answer']:
            sub_exercises = Exercise.objects.filter(
                parent_id=exercise.id
            ).order_by('order').values(
                'id', 'order', 'content', 'options', 'question_type', 'answer'
            )

        context.update({
            'exercise': exercise,
            'knowledge_points': knowledge_points,
            'methods': methods,
            'answer_records': answer_records,
            'sub_exercises': list(sub_exercises),
            'show_answer_records': exercise.question_type not in ['cloze_test', 'reading_answer']
        })
        return context

        
    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            # 使用prefetch_related优化关联知识点查询
            exercise = self.get_object()
            exercise = Exercise.objects.prefetch_related(
                'exercise_knowledge_points__knowledge_point'
            ).get(id=exercise.id)
            
            # 获取关联知识点信息
            knowledge_points = []
            for rel in exercise.exercise_knowledge_points.all():
                knowledge_points.append({
                    'id': rel.knowledge_point.id,
                    'title': rel.knowledge_point.title
                })
            
            return JsonResponse({
                'id': exercise.id,
                'title': exercise.title,
                'question_type': exercise.question_type,
                'content': exercise.content,
                'options': exercise.options,
                'answer': exercise.answer,
                'analysis': exercise.analysis,
                'difficulty': exercise.difficulty,
                'memory_level': exercise.memory_level,
                'mastery_level': exercise.mastery_level,
                'parent_id': exercise.parent_id,
                'chapter': exercise.chapter_id,
                'knowledge_points': knowledge_points  # 新增知识点信息
            })
        return super().get(request, *args, **kwargs)

def submit_answer(request, pk):
    if request.method == 'POST':
        exercise = Exercise.objects.get(pk=pk)
        answer = request.POST.get('answer', '')
        is_correct = (answer.strip().lower() == exercise.answer.strip().lower())
        
        # 保存答题记录
        submission = exercise.answers.create(
            user=request.user,
            answer=answer,
            is_correct=is_correct
        )
        
        return JsonResponse({
            'success': True,
            'is_correct': is_correct,
            'correct_answer': exercise.answer,
            'analysis': exercise.analysis,
            'created_at': submission.created_at.strftime('%Y-%m-%d %H:%M')
        })
    return JsonResponse({'success': False}, status=400)

def delete_exercise(request, pk):
    if request.method == 'POST':
        exercise = Exercise.objects.get(pk=pk)
        exercise.delete()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False}, status=400)

# 答题相关视图类
class ExerciseSetStartAPI(generics.RetrieveAPIView):
    """开始答题，返回第一道习题"""
    serializer_class = ExerciseSerializer
    
    def get_object(self):
        exercise_set = get_object_or_404(ExerciseSet, id=self.kwargs['exercise_set_id'])
        exercises = exercise_set.exercises.all().order_by('id')
        
        if not exercises.exists():
            return None
            
        first_exercise = exercises.first()
        
        # 如果是完形填空或阅读理解，返回主问题和所有子问题
        if first_exercise.question_type in ['cloze_test', 'reading_answer']:
            sub_questions = Exercise.objects.filter(parent_id=first_exercise.id).order_by('id')
            if not sub_questions.exists():
                return first_exercise
            return {
                'main_question': first_exercise,
                'sub_questions': sub_questions
            }
        return first_exercise

    def retrieve(self, request, *args, **kwargs):
        exercise_set = get_object_or_404(ExerciseSet, id=self.kwargs['exercise_set_id'])
        
        # 如果是API请求，返回JSON数据
        if request.accepted_renderer.format == 'json':
            exercise = self.get_object()
            if not exercise:
                return Response(
                    {'error': '练习集中没有习题'}, 
                    status=status.HTTP_404_NOT_FOUND
                )
                
            if isinstance(exercise, dict):  # 处理完形填空/阅读理解
                serializer = self.get_serializer(exercise['main_question'])
                sub_serializer = self.get_serializer(exercise['sub_questions'], many=True)
                return Response({
                    'main_question': serializer.data,
                    'sub_questions': sub_serializer.data
                })
                
            serializer = self.get_serializer(exercise)
            return Response(serializer.data)
        
        # 否则返回答题页面
        exercise_count = exercise_set.exercises.count()
        if exercise_count == 0:
            return Response(
                {'error': '练习集中没有习题'}, 
                status=status.HTTP_404_NOT_FOUND
            )
            
        from django.shortcuts import render
        return render(request, 'courses/exercise_set_quiz.html', {
            'exercise_set': exercise_set,
            'exercise_count': exercise_count
        })
class ExerciseSetStartQuickAPI(generics.RetrieveAPIView):
    """开始答题，返回第一道习题"""
    serializer_class = ExerciseSerializer
    
    def get_object(self):
        exercise_set = get_object_or_404(ExerciseSet, id=self.kwargs['exercise_set_id'])
        exercises = exercise_set.exercises.all().order_by('id')
        
        if not exercises.exists():
            return None
            
        first_exercise = exercises.first()
        
        # 如果是完形填空或阅读理解，返回主问题和所有子问题
        if first_exercise.question_type in ['cloze_test', 'reading_answer']:
            sub_questions = Exercise.objects.filter(parent_id=first_exercise.id).order_by('id')
            if not sub_questions.exists():
                return first_exercise
            return {
                'main_question': first_exercise,
                'sub_questions': sub_questions
            }
        return first_exercise

    def retrieve(self, request, *args, **kwargs):
        exercise_set = get_object_or_404(ExerciseSet, id=self.kwargs['exercise_set_id'])
        exercise_set_id = self.kwargs['exercise_set_id']
        
        # 如果是API请求，返回JSON数据
        if request.accepted_renderer.format == 'json':
            exercise = self.get_object()
            if not exercise:
                return Response(
                    {'error': '练习集中没有习题'}, 
                    status=status.HTTP_404_NOT_FOUND
                )
                
            if isinstance(exercise, dict):  # 处理完形填空/阅读理解
                serializer = self.get_serializer(exercise['main_question'])
                sub_serializer = self.get_serializer(exercise['sub_questions'], many=True)
                return Response({
                    'main_question': serializer.data,
                    'sub_questions': sub_serializer.data
                })
                
            serializer = self.get_serializer(exercise)
            return Response(serializer.data)
        
        # 否则返回答题页面
        exercise_count = exercise_set.exercises.count()
        if exercise_count == 0:
            return Response(
                {'error': '练习集中没有习题'}, 
                status=status.HTTP_404_NOT_FOUND
            )
            
        from django.shortcuts import render
        return render(request, 'courses/exercise_set_quiz_quick.html', {
            'exercise_set': exercise_set,
            'exercise_count': exercise_count,
            'exercise_set_id': exercise_set_id
        })

class ExerciseSetNextAPI(generics.RetrieveAPIView):
    """获取下一道习题"""
    serializer_class = ExerciseSerializer
    
    def get_object(self):
        exercise_set = get_object_or_404(ExerciseSet, id=self.kwargs['exercise_set_id'])
        current_exercise = get_object_or_404(Exercise, id=self.kwargs['current_exercise_id'])
        
        # 获取当前习题在练习集中的位置
        exercises = list(exercise_set.exercises.order_by('id'))
        try:
            current_index = exercises.index(current_exercise)
            next_exercise = exercises[current_index + 1]
            
            # 处理完形填空/阅读理解
            if next_exercise.question_type in ['cloze_test', 'reading_answer']:
                return {
                    'main_question': next_exercise,
                    'sub_questions': Exercise.objects.filter(parent_id=next_exercise.id).order_by('id')
                }
            return next_exercise
        except (ValueError, IndexError):
            return None  # 没有下一题了

    def retrieve(self, request, *args, **kwargs):
        exercise = self.get_object()
        if not exercise:
            return Response(
                {'message': '已经是最后一道题了'}, 
                status=status.HTTP_404_NOT_FOUND
            )
            
        if isinstance(exercise, dict):  # 处理完形填空/阅读理解
            serializer = self.get_serializer(exercise['main_question'])
            sub_serializer = self.get_serializer(exercise['sub_questions'], many=True)
            return Response({
                'main_question': serializer.data,
                'sub_questions': sub_serializer.data
            })
            
        serializer = self.get_serializer(exercise)
        return Response(serializer.data)

class ExerciseNeighborAPI(APIView):
    """
    获取相邻习题API
    GET /api/exercises/<exercise_id>/neighbor/?direction=prev|next&chapter_id=<chapter_id>
    """
    def get(self, request, exercise_id):
        direction = request.GET.get('direction', 'next')
        chapter_id = request.GET.get('chapter_id')
        
        if direction not in ['prev', 'next']:
            return Response(
                {'error': 'direction参数必须是prev或next'}, 
                status=status.HTTP_400_BAD_REQUEST
            )
        
        try:
            current_exercise = Exercise.objects.get(id=exercise_id)
            chapter = Chapter.objects.get(id=chapter_id) if chapter_id else current_exercise.chapter
            
            # 获取同一章节的所有习题(按创建时间排序)
            exercises = Exercise.objects.filter(
                chapter=chapter,
                parent_id__isnull=True  # 只查询主习题
            ).order_by('created_at')
            
            # 查找当前习题在列表中的位置
            exercise_list = list(exercises)
            try:
                current_index = exercise_list.index(current_exercise)
                
                if direction == 'prev' and current_index > 0:
                    neighbor = exercise_list[current_index - 1]
                elif direction == 'next' and current_index < len(exercise_list) - 1:
                    neighbor = exercise_list[current_index + 1]
                else:
                    if direction == 'prev':
                        return Response(
                            {'message': '没有上一题了'}, 
                            status=status.HTTP_404_NOT_FOUND
                        )
                    else:   
                        return Response(
                            {'message': '没有下一题了'},
                            status=status.HTTP_404_NOT_FOUND
                    )
                
                serializer = ExerciseSerializer(neighbor)
                return Response(serializer.data)
                
            except ValueError:
                return Response(
                    {'error': '当前习题不在指定章节中'},
                    status=status.HTTP_400_BAD_REQUEST
                )
                
        except Exercise.DoesNotExist:
            return Response(
                {'error': '习题不存在'}, 
                status=status.HTTP_404_NOT_FOUND
            )
        except Chapter.DoesNotExist:
            return Response(
                {'error': '章节不存在'}, 
                status=status.HTTP_404_NOT_FOUND
            )

class ExerciseSetSubmitAPI(generics.CreateAPIView):
    """提交练习集并创建完成记录"""
    serializer_class = ExerciseSetCompletionSerializer

    def create(self, request, *args, **kwargs):
        try:
            exercise_set = get_object_or_404(ExerciseSet, pk=kwargs['exercise_set_id'])
            user = request.user
            answers_data = request.data.get('answers', {})
            total_time = int(request.data.get('total_time', 0))

            answer_records = []
            correct_count = 0
            
            from django.db import transaction
            with transaction.atomic():
                # 1. 先创建临时练习集完成记录(correct_count=0)
                completion = ExerciseSetCompletion.objects.create(
                    user=user,
                    exercise_set=exercise_set,
                    score=0,
                    correct_count=0,
                    total_count=exercise_set.exercises.count(),
                    time_spent=total_time
                )

                # 2. 创建所有习题答题记录并计算正确数量
                correct_count = 0
                for exercise in exercise_set.exercises.all():
                    exercise_id_str = str(exercise.id)  # 统一转换为字符串
                    answer_info = answers_data.get(exercise_id_str, {})
                    user_answer = answer_info.get('answer', '')
                    time_spent = answer_info.get('time_spent', 0)
                    difficulty = answer_info.get('difficulty', 0)
                    memory_level = answer_info.get('memory_level', 0)
                    mastery_level = answer_info.get('mastery_level', 0)
                    exercise_id = answers_data.get(exercise.id)
                    
                    # 处理主习题
                    is_correct = str(user_answer) == str(exercise.answer)
                    if is_correct:
                        correct_count += 1
                    
                    answer_record = ExerciseAnswer.objects.create(
                        exercise=exercise,
                        user=user,
                        exercise_set_completion=completion,
                        answer=user_answer,
                        is_correct=is_correct,
                        time_spent=time_spent,
                        difficulty=difficulty,
                        memory_level=memory_level,
                        mastery_level=mastery_level,
                    )
                    answer_records.append({
                        'exercise_id': exercise.id,
                        'answer_id': answer_record.id,
                        'is_correct': is_correct
                    })
                    
                    # 处理子习题（完形填空/阅读理解）
                    if exercise.question_type in ['cloze_test', 'reading_answer']:
                        for sub_exercise in Exercise.objects.filter(parent_id=exercise.id):
                            sub_answer_info = user_answer.get(str(sub_exercise.id), {})
                            sub_answer = sub_answer_info.get('answer', '')
                            sub_time_spent = sub_answer_info.get('time_spent', 0)
                            
                            sub_is_correct = str(sub_answer) == str(sub_exercise.answer)
                            if sub_is_correct:
                                correct_count += 1
                            
                            sub_record = ExerciseAnswer.objects.create(
                                exercise=sub_exercise,
                                user=user,
                                exercise_set_completion=completion,
                                answer=sub_answer,
                                is_correct=sub_is_correct,
                                time_spent=sub_time_spent
                            )
                            answer_records.append({
                                'exercise_id': sub_exercise.id,
                                'answer_id': sub_record.id,
                                'is_correct': sub_is_correct
                            })

                # 3. 更新练习集完成记录的正确数量和分数
                completion.correct_count = correct_count
                completion.score = (correct_count / exercise_set.exercises.count()) * 100 if exercise_set.exercises.count() else 0
                completion.save()


            return Response({
                'redirect_url': reverse('courses:exercise-set-result', kwargs={
                    'exercise_set_id': exercise_set.id,
                    'completion_id': completion.id
                }),
                'completion_id': completion.id,
                'message': '提交成功'
            }, status=status.HTTP_201_CREATED)

        except Exception as e:
            return Response(
                {'error': str(e)},
                status=status.HTTP_400_BAD_REQUEST
            )
    
    def calculate_score(self, data, exercise_set):
        correct = 0
        answers = data.get('answers', {})
        for exercise in exercise_set.exercises.all():
            if str(answers.get(str(exercise.id))) == str(exercise.answer):
                correct += 1
        return (correct / exercise_set.exercises.count()) * 100 if exercise_set.exercises.count() else 0
    """提交练习集并创建完成记录"""


# 原有习题操作函数
@api_view(['POST'])
def submit_answer(request, pk):
    """提交单个习题答案"""

    
    try:
        exercise = get_object_or_404(Exercise, id=pk)
        user_answer = request.data.get('user_answer', '').strip()
        
        if not user_answer:
            return Response(
                {'error': '答案不能为空'}, 
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # 验证答案是否正确
        is_correct = str(user_answer).lower() == str(exercise.answer).lower()
        
        # 创建答题记录
        answer_record = ExerciseAnswer.objects.create(
            exercise=exercise,
            user=request.user,
            answer=user_answer,
            is_correct=is_correct,
            time_spent=request.data.get('time_spent', 0),
            difficulty=request.data.get('difficulty', 0),
            memory_level=request.data.get('memory_level', 0),
            mastery_level=request.data.get('mastery_level', 0),
            created_at=timezone.now()
        )
        
        # 更新习题统计数据
        if is_correct:
            exercise.correct_count = (exercise.correct_count or 0) + 1
        else:
            exercise.wrong_count = (exercise.wrong_count or 0) + 1
        exercise.save()
        
        # 返回响应
        response_data = {
            'status': 'success',
            'is_correct': is_correct,
            'correct_answer': exercise.answer,
            'analysis': exercise.analysis,
            'answer_id': answer_record.id,
            'created_at': answer_record.created_at.strftime('%Y-%m-%d %H:%M:%S')
        }
        
        return Response(response_data, status=status.HTTP_201_CREATED)
        
    except Exception as e:
        return Response(
            {'error': str(e)}, 
            status=status.HTTP_500_INTERNAL_SERVER_ERROR
        )

def exercise_set_exercise_delete(request, exercise_set_id, exercise_id):
    """从习题集中删除习题(仅删除关联关系)"""
    exercise_set = get_object_or_404(ExerciseSet, id=exercise_set_id)
    exercise = get_object_or_404(Exercise, id=exercise_id)
    
    # 从多对多关系中移除
    exercise_set.exercises.remove(exercise)
    
    return JsonResponse({
        'success': True,
        'message': '已从习题集中移除该习题',
        'remaining_count': exercise_set.exercises.count()
    })

def exercise_set_exercise_add(request, exercise_set_id):
    """向习题集中添加习题"""
    exercise_set = get_object_or_404(ExerciseSet, id=exercise_set_id)
    
    if request.method == 'POST':
        try:
            # 从JSON请求体中获取数据
            data = json.loads(request.body)
            exercise_ids = data.get('exercise_ids', [])
            
            if not exercise_ids:
                return JsonResponse({
                    'success': False,
                    'message': '请提供要添加的习题ID'
                }, status=400)
                
            # 获取所有有效的习题
            exercises = Exercise.objects.filter(id__in=exercise_ids)
            if exercises.count() != len(exercise_ids):
                return JsonResponse({
                    'success': False,
                    'message': '部分习题ID无效'
                }, status=400)
                
            # 批量添加到习题集
            exercise_set.exercises.add(*exercises)
            
            return JsonResponse({
                'success': True,
                'message': f'成功添加 {len(exercise_ids)} 个习题到习题集',
                'new_count': exercise_set.exercises.count()
            })
            
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'数据解析错误: {str(e)}'
            }, status=400)
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'添加失败: {str(e)}'
            }, status=400)
    
    return JsonResponse({
        'success': False,
        'message': '无效的请求方法'
    }, status=405)

def get_exercise_edit(request, pk):
    """获取习题编辑数据"""
    exercise = get_object_or_404(Exercise, id=pk)
    
    if request.method == 'GET':
        serializer = ExerciseSerializer(exercise)
        return JsonResponse(serializer.data)
    
    return JsonResponse({'success': False}, status=400)

def update_exercise_edit(request, pk):
    if request.method == 'POST':
        try:
            # 获取表单数据
            data = request.POST.dict()
            
            # 处理options字段
            if 'options' in data:
                try:
                    data['options'] = json.loads(data['options'])
                except json.JSONDecodeError as e:
                    return JsonResponse({
                        'success': False,
                        'message': 'options字段必须是有效的JSON格式'
                    }, status=400)
            
            # 验证必要字段
            required_fields = ['title', 'question_type', 'answer']
            for field in required_fields:
                if field not in data or not data[field]:
                    return JsonResponse({
                        'success': False,
                        'message': f'缺少必要字段: {field}'
                    }, status=400)
            
            # 获取并更新Exercise对象
            try:
                exercise = Exercise.objects.get(pk=pk)
                
                # 基础字段更新
                exercise.title = data.get('title', exercise.title)
                exercise.question_type = data.get('question_type', exercise.question_type)
                exercise.content = data.get('content', exercise.content)
                exercise.answer = data.get('answer', exercise.answer)
                exercise.analysis = data.get('analysis', exercise.analysis)
                
                # 处理options字段
                if 'options' in data:
                    exercise.options = data['options']
                
                # 数值型字段处理
                try:
                    exercise.difficulty = int(data.get('difficulty', exercise.difficulty))
                    exercise.memory_level = int(data.get('memory_level', exercise.memory_level))
                    exercise.mastery_level = int(data.get('mastery_level', exercise.mastery_level))
                    exercise.answer_time = int(data.get('answer_time', exercise.answer_time))
                    exercise.wrong_count = int(data.get('wrong_count', exercise.wrong_count))
                except (ValueError, TypeError):
                    pass  # 保持原值
                
                # 关联字段处理
                if 'chapter' in data:
                    try:
                        exercise.chapter = Chapter.objects.get(pk=int(data['chapter']))
                    except (Chapter.DoesNotExist, ValueError):
                        pass
                
                # 日期字段处理-无修改不处理

                exercise.save()
                return JsonResponse({'success': True})
                
            except Exercise.DoesNotExist:
                return JsonResponse({
                    'success': False,
                    'message': '习题不存在'
                }, status=404)


            return JsonResponse({'success': True})
            
        except Exception as e:
            return JsonResponse({
                'success': False,
                'message': f'服务器错误: {str(e)}'
            }, status=500)
def get_exercise_parentid(request, pk):
    if request.method == 'GET':
        try:
            exercises = Exercise.objects.filter(parent_id=pk).order_by('order').values(
                'id',
                'question_type', 
                'content',
                'options',
                'answer',
                'order'
            )
            return JsonResponse({
                'exercises': list(exercises),
                'count': len(exercises)
            }, safe=False)
            
        except Exercise.DoesNotExist:
            return JsonResponse({'error': '练习不存在'}, status=404)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)
            
    return JsonResponse({'error': 'Invalid method'}, status=405)
    